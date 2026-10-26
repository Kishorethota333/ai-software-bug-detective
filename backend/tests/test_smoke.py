import io
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from app.main import app

def test_health():
    client = TestClient(app)
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"

def test_core_workflow():
    import uuid
    client = TestClient(app)
    email = f"test-{uuid.uuid4().hex[:10]}@example.com"
    registered = client.post("/api/auth/register", json={"name":"Test Developer", "email":email, "password":"correct-horse-battery"})
    assert registered.status_code == 201, registered.text
    token = registered.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    project = client.post("/api/projects", headers=headers, json={"name":"Test Project", "description":"test", "technology":"FastAPI"})
    assert project.status_code == 201, project.text
    bug = client.post("/api/bugs", headers=headers, json={"title":"API returns 500 error", "project_id":project.json()["id"], "error_message":"Internal Server Error 500"})
    assert bug.status_code == 201, bug.text
    analysis = client.post(f"/api/bugs/{bug.json()['id']}/analyze", headers=headers)
    assert analysis.status_code == 200, analysis.text
    assert analysis.json()["root_cause"]
    cases = client.get("/api/test-cases", headers=headers, params={"bug_id":bug.json()["id"]})
    assert cases.status_code == 200 and len(cases.json()) >= 1
    workbook_response = client.get("/api/test-cases/excel", headers=headers)
    assert workbook_response.status_code == 200
    workbook = load_workbook(io.BytesIO(workbook_response.content), read_only=True)
    sheet = workbook["Test Cases"]
    all_cases_response = client.get("/api/test-cases", headers=headers)
    all_cases = all_cases_response.json()
    assert sheet.max_row == len(all_cases) + 1
    assert [cell.value for cell in sheet[1]] == [
        "Test Case ID",
        "Bug ID",
        "Scenario",
        "Input",
        "Expected Result",
        "Type",
        "Priority",
        "Marked Executed",
    ]
    exported_scenarios = {
        row[2]
        for row in sheet.iter_rows(min_row=2, values_only=True)
    }
    assert exported_scenarios == {case["scenario"] for case in all_cases}
    workbook.close()
